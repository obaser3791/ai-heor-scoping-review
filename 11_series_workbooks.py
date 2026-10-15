"""Per article workbooks for the two part series, with renumbered table ids."""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = "/home/user/workspace"
SPEC = json.load(open(os.path.join(ROOT, "tables", "tables_spec.json")))
BY = {t["id"]: t for t in SPEC}

HFILL = PatternFill("solid", fgColor="1B474D")
TFILL = PatternFill("solid", fgColor="F4F4F2")
_s = Side(style="thin", color="D4D1CA")
BORD = Border(left=_s, right=_s, top=_s, bottom=_s)


def remap(pairs, xrefs=()):
    out = []
    for old, new in pairs:
        t = json.loads(json.dumps(BY[old]))
        t["id"] = new
        if t.get("footnote"):
            for a, b in xrefs:
                t["footnote"] = t["footnote"].replace(a, b)
        out.append(t)
    return out


def write_wb(tabs, path):
    wb = Workbook()
    wb.remove(wb.active)
    for t in tabs:
        ws = wb.create_sheet(t["id"].replace(" ", ""))
        ws["A1"] = f"{t['id']}. {t['title']}"
        ws["A1"].font = Font(bold=True, size=12, color="1B474D")
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(2, len(t["headers"])))
        ws.row_dimensions[1].height = 34
        for j, h in enumerate(t["headers"], start=1):
            c = ws.cell(row=3, column=j, value=str(h))
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = HFILL
            c.border = BORD
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[3].height = 32
        for i, row in enumerate(t["rows"], start=4):
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(size=9, bold=(j == 1 and bool(v) and not str(v).startswith("  ")
                                            and len(t["headers"]) == 4))
                c.border = BORD
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if i % 2 == 0:
                for j in range(1, len(t["headers"]) + 1):
                    ws.cell(row=i, column=j).fill = TFILL
        fr = len(t["rows"]) + 5
        ws.cell(row=fr, column=1, value="Note. " + (t.get("footnote") or "")).font = Font(size=8, italic=True)
        ws.cell(row=fr, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fr, start_column=1, end_row=fr,
                       end_column=max(2, len(t["headers"])))
        ws.row_dimensions[fr].height = 60
        for j in range(1, len(t["headers"]) + 1):
            col = ws.cell(row=3, column=j).column_letter
            vals = [len(str(r[j - 1])) for r in t["rows"] if j - 1 < len(r)]
            mx = max(vals) if vals else 10
            ws.column_dimensions[col].width = min(60, max(12, mx * 0.55 if mx > 40 else mx + 4))
        ws.freeze_panes = "A4"
    wb.save(path)
    return path


P2X = [("eTable 15", "eTable 1"), ("eTable 16", "eTable 2"), ("eTable 17", "eTable 3"),
       ("eTable 18", "eTable 4"), ("eTable 19", "eTable 5")]

jobs = [
    (remap([("Table 1", "Table 1"), ("Table 2", "Table 2")]), "part1_main_tables.xlsx"),
    (remap([(f"eTable {k}", f"eTable {k}") for k in range(1, 15)]), "part1_supplementary_tables.xlsx"),
    (remap([("Table 3", "Table 1"), ("Table 4", "Table 2"), ("Table 5", "Table 3")]), "part2_main_tables.xlsx"),
    (remap(P2X, P2X), "part2_supplementary_tables.xlsx"),
]

outdir = os.path.join(ROOT, "tables", "series")
os.makedirs(outdir, exist_ok=True)
for tabs, name in jobs:
    p = write_wb(tabs, os.path.join(outdir, name))
    print("wrote", p, len(tabs), "sheets")
